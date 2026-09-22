from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\ulab-kalibrasi")
DOCS = ROOT / "docs"
ASSETS = DOCS / "panduan-temuan-assets"
OUTPUT = DOCS / "Panduan_Penggunaan_Fitur_Temuan_Ketidaksesuaian.docx"

NAVY = "0B2545"
INK = "1F2937"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
MUTED = "667085"
LIGHT_BLUE = "E9F7FB"
LIGHT_GRAY = "F4F6F9"
LIGHT_GOLD = "FFF4DD"
LIGHT_RED = "FDECEC"
GREEN = "18A878"


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D8DEE6", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def configure_style(style, size, color=INK, bold=False, before=0, after=6, line=1.25):
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_begin, instr, fld_end))


def add_numbering_definition(doc, kind="decimal"):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if kind == "decimal" else "bullet")
    lvl.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1." if kind == "decimal" else "•")
    lvl.append(text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, nid))


def add_numbered(doc, items):
    num_id = add_numbering_definition(doc, "decimal")
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        apply_num(p, num_id)


def add_bullets(doc, items):
    num_id = add_numbering_definition(doc, "bullet")
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        apply_num(p, num_id)


def add_callout(doc, label, text, fill=LIGHT_BLUE, border="7EC8DA"):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.right_indent = Pt(8)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(f"{label}  ")
    r.bold = True
    p.add_run(text)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "5")
        el.set(qn("w:color"), border)
        p_bdr.append(el)
    p_pr.append(p_bdr)
    return p


def add_figure(doc, filename, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    picture = run.add_picture(str(ASSETS / filename), width=Inches(6.4))
    doc_pr = picture._inline.docPr
    doc_pr.set("descr", caption)
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(10)
    for r in cap.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = RGBColor.from_string(MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_access_table(doc):
    rows = [
        ("Peran", "Yang terlihat", "Yang dapat dilakukan"),
        ("Auditor", "Lingkup tempat akun ditugaskan sebagai Auditor", "Tambah, lihat, edit, dan hapus temuan pada lingkupnya"),
        ("Auditee", "Lingkup tempat akun ditugaskan sebagai Auditee", "Melihat temuan pada lingkupnya (lihat saja)"),
        ("Lead Auditor", "Ditampilkan satu kali sebagai ketua Tim Audit Internal", "Mendukung seluruh lingkup sesuai penetapan tahun"),
        ("Pengelola", "Mapping Auditor Internal per tahun", "Menetapkan anggota, Auditee, status Pakta, dan mencetak Pakta"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    set_table_geometry(table, [1800, 3670, 3890], 120)
    set_repeat_table_header(table.rows[0])
    for i, row_data in enumerate(rows):
        for j, value in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if i == 0:
                shade_cell(cell, "E8EEF5")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(10 if i else 10.5)
                    r.font.bold = i == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def body_paragraph(doc, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=INK, after=6):
    p = doc.add_paragraph(style="Normal")
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def build():
    DOCS.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        OUTPUT.unlink()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    configure_style(doc.styles["Normal"], 11, INK, False, 0, 6, 1.25)
    configure_style(doc.styles["Title"], 30, NAVY, True, 0, 8, 1.0)
    configure_style(doc.styles["Subtitle"], 14, "52677D", False, 0, 8, 1.0)
    configure_style(doc.styles["Heading 1"], 16, BLUE, True, 18, 10, 1.0)
    configure_style(doc.styles["Heading 2"], 13, BLUE, True, 14, 7, 1.0)
    configure_style(doc.styles["Heading 3"], 12, BLUE_DARK, True, 10, 5, 1.0)

    header = section.header.paragraphs[0]
    header.text = "PANDUAN FITUR TEMUAN KETIDAKSESUAIAN"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("ULAB | Panduan Pengguna Fitur Temuan Ketidaksesuaian")
    for run in footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    first_footer = section.first_page_footer.paragraphs[0]
    first_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first_footer.add_run("ULAB | Panduan Pengguna | 01 September 2026")
    for run in first_footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    for _ in range(4):
        doc.add_paragraph()
    kicker = body_paragraph(doc, "PANDUAN PENGGUNA", True, WD_ALIGN_PARAGRAPH.CENTER, GREEN, 4)
    kicker.runs[0].font.size = Pt(11)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Fitur Temuan Ketidaksesuaian")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Panduan praktis untuk Auditor Internal, Auditee, dan Pengelola Mapping")
    body_paragraph(doc, "FMMO-163-14.4.3.b-88.5", True, WD_ALIGN_PARAGRAPH.CENTER, NAVY, 2)
    body_paragraph(doc, "Laporan Ringkas dan Lembar Temuan Ketidaksesuaian", False, WD_ALIGN_PARAGRAPH.CENTER, MUTED, 18)
    add_callout(doc, "INTI PENGGUNAAN", "Setelah login, sistem langsung mengikuti mapping tahunan. Auditor mengisi temuan pada lingkup tugasnya; Auditee melihat hasil pada lingkup tempat ia dipetakan sebagai Auditee.")
    body_paragraph(doc, "Versi panduan: 01 September 2026", False, WD_ALIGN_PARAGRAPH.CENTER, MUTED, 0)

    add_page_break(doc)
    add_heading(doc, "1. Ringkasan Alur")
    body_paragraph(doc, "Gunakan alur berikut sebagai gambaran cepat sebelum membaca langkah rinci.")
    add_figure(doc, "00-alur-ringkas.png", "Gambar 1. Alur penggunaan dari login sampai pencetakan.")
    add_callout(doc, "PENTING", "Jenis audit tidak dipilih manual oleh pengguna. Lingkup yang muncul ditentukan oleh akun login, mapping peran, dan tahun laporan yang dipilih.", LIGHT_GOLD, "F5A623")

    add_page_break(doc)
    add_heading(doc, "2. Masuk ke Fitur Setelah Login")
    add_numbered(doc, [
        "Login ke aplikasi ULAB dengan akun masing-masing.",
        "Klik NAVIGATION pada bagian atas aplikasi.",
        "Pada kelompok PROGRAM MUTU, klik Temuan Ketidaksesuaian.",
    ])
    add_figure(doc, "01-navigation.png", "Gambar 2. Jalur menu: Navigation -> Program Mutu -> Temuan Ketidaksesuaian.")
    add_callout(doc, "CATATAN PENGELOLA", "Menu Mapping Auditor Internal digunakan untuk menyiapkan penugasan per tahun. Auditor dan Auditee biasa tidak perlu memilih perannya sendiri.", LIGHT_GRAY, "B9C2CF")

    add_page_break(doc)
    add_heading(doc, "3. Peran Sudah Dimapping Sesuai Tugas")
    body_paragraph(doc, "Sesudah halaman terbuka, sistem membaca mapping pada tahun yang dipilih. Satu akun dapat menjadi Auditor pada suatu lingkup dan sekaligus menjadi Auditee pada lingkup lain.")
    add_access_table(doc)
    add_callout(doc, "URUTAN TAMPILAN", "Lingkup tempat pengguna menjadi Auditor selalu ditampilkan lebih dahulu. Lingkup tempat pengguna menjadi Auditee ditampilkan di bawah dengan label Mode Auditee | Lihat Saja.")
    add_figure(doc, "04-peran-auditor-auditee.png", "Gambar 3. Perbedaan tampilan Auditor (atas) dan Auditee (bawah).")

    add_page_break(doc)
    add_heading(doc, "4. Kenali Tombol Utama dan Periode Laporan")
    body_paragraph(doc, "Pastikan tahun pada Periode Laporan sesuai dengan tahun audit. Gunakan Muat Ulang setelah mengganti tahun agar mapping dan data ditampilkan kembali.")
    add_figure(doc, "02-toolbar.png", "Gambar 4. Tombol utama Pakta, pencetakan, dan label peran akun.")
    add_numbered(doc, [
        "Pakta Integritas: membuka atau mengisi Pakta Auditor untuk akun yang sedang login.",
        "Cetak Pakta: mencetak Pakta milik akun login setelah Pakta tersimpan.",
        "Cetak Laporan: membuka satu laporan gabungan seluruh lingkup pada tahun terpilih.",
    ])
    add_callout(doc, "SYARAT CETAK LAPORAN", "Tombol Cetak Laporan aktif setelah minimal satu data temuan tersimpan pada tahun yang dipilih, termasuk temuan yang dibuat oleh Auditor lain.", LIGHT_GOLD, "F5A623")

    add_page_break(doc)
    add_heading(doc, "5. Auditor: Isi Pakta Integritas Terlebih Dahulu")
    body_paragraph(doc, "Pakta Integritas wajib bagi Auditor sebelum mengelola temuan. Auditee murni tidak wajib mengisi Pakta Auditor.")
    add_figure(doc, "03-pakta.png", "Gambar 5. Form Pakta Integritas dan area tanda tangan digital.")
    add_numbered(doc, [
        "Periksa Nama, NID, dan Jabatan Tim Audit Internal. Data ini terisi otomatis dari akun login dan mapping.",
        "Isi Tanggal Pernyataan dan pilih Lokasi Jakarta atau Gresik.",
        "Bubuhkan tanda tangan pada area tanda tangan menggunakan mouse, stylus, atau layar sentuh.",
        "Klik Simpan Pakta Integritas.",
    ])
    add_callout(doc, "JANGAN TERTUKAR", "Jabatan yang tampil pada Pakta adalah jabatan di Tim Audit Internal (misalnya Anggota Auditor atau Auditor Observer), bukan jabatan struktural pegawai.", LIGHT_RED, "E76A6A")

    add_page_break(doc)
    add_heading(doc, "6. Auditor: Tambah Temuan Ketidaksesuaian")
    body_paragraph(doc, "Pada kartu lingkup tempat Anda menjadi Auditor, klik Tambah Temuan. Tombol ini tidak muncul pada kartu Mode Auditee.")
    add_figure(doc, "05-tambah-temuan.png", "Gambar 6. Form tambah temuan; bagian dan nama pengisi mengikuti mapping/login.")
    add_numbered(doc, [
        "Isi Klausul sesuai standar acuan dan pilih Kategori Temuan: 1 - Major, 2 - Minor, atau 3 - Observasi.",
        "Tuliskan uraian ketidaksesuaian secara lengkap: kondisi, bukti objektif, dan dokumen/rekaman terkait.",
        "Pastikan Auditor/Pengisi sudah menunjukkan nama akun yang sedang login.",
        "Klik Simpan Temuan.",
    ])
    add_bullets(doc, [
        "Nama LPK dan Standar Acuan diisi otomatis oleh sistem.",
        "Tanggal Audit Internal ditetapkan otomatis saat temuan pertama pada lingkup tersebut disimpan.",
        "Data yang baru disimpan langsung menambah ringkasan Major, Minor, Observasi, dan Total.",
    ])

    add_page_break(doc)
    add_heading(doc, "7. Auditee: Melihat Temuan yang Ditujukan kepada Anda")
    body_paragraph(doc, "Jika akun Anda juga dipetakan sebagai Auditee, kartu lingkup tersebut muncul di bawah bagian Auditor dan diberi label Mode Auditee | Lihat Saja.")
    add_numbered(doc, [
        "Cari kartu dengan label Mode Auditee | Lihat Saja.",
        "Lihat daftar temuan, bagian, klausul, kategori, uraian, serta Auditor/Pengisi.",
        "Tidak ada tombol Tambah Temuan, Edit, atau Hapus pada mode Auditee.",
    ])
    add_callout(doc, "CONTOH", "Seseorang dapat menjadi Auditor Kelistrikan Jakarta dan sekaligus menjadi Auditee Kelistrikan Gresik. Temuan sebagai Auditor tampil di bagian atas; temuan yang ditujukan kepadanya sebagai Auditee tampil di bagian bawah.")
    add_heading(doc, "Batas akses yang perlu dipahami", 2)
    add_bullets(doc, [
        "Auditor hanya mengelola temuan pada lingkup tempat ia ditugaskan sebagai Auditor.",
        "Auditee dapat melihat seluruh temuan pada lingkup tempat ia ditetapkan sebagai Auditee.",
        "Pengaturan mengikuti mapping tahun dan akun login, bukan siapa yang membuat data lebih dahulu.",
    ])

    add_page_break(doc)
    add_heading(doc, "8. Melihat Mapping dan Mencetak Pakta")
    add_figure(doc, "06-mapping.png", "Gambar 7. Mapping tahunan, status pengisian Pakta, dan ikon cetak per Auditor.")
    add_heading(doc, "Untuk pengguna", 2)
    add_bullets(doc, [
        "Klik Cetak Pakta pada halaman Temuan Ketidaksesuaian untuk mencetak Pakta milik akun login.",
        "Cetak Pakta hanya tersedia setelah Pakta Integritas berhasil disimpan.",
    ])
    add_heading(doc, "Untuk pengelola mapping", 2)
    add_bullets(doc, [
        "Lead Auditor ditampilkan satu kali dan tidak perlu diulang pada setiap kartu lingkup.",
        "Setiap kartu lingkup berisi anggota Auditor dan Auditee sesuai penugasan.",
        "Status Pakta sudah diisi/belum diisi terlihat pada baris Auditor.",
        "Klik ikon printer pada Auditor yang sudah mengisi Pakta untuk membuka cetakan Pakta orang tersebut.",
        "Mapping diperiodekan per tahun sehingga dapat disiapkan kembali untuk tahun berikutnya.",
    ])

    add_page_break(doc)
    add_heading(doc, "9. Mencetak Laporan Keseluruhan")
    body_paragraph(doc, "Klik Cetak Laporan untuk melihat hasil keseluruhan pada tahun yang dipilih. Cetakan tidak dibatasi oleh lingkup yang sedang terlihat pada akun login.")
    add_figure(doc, "07-cetak-laporan.png", "Gambar 8. Cetak Laporan menghasilkan satu PDF gabungan seluruh lingkup.")
    add_bullets(doc, [
        "Semua lingkup yang sudah memiliki data temuan digabung dalam satu PDF.",
        "Isi PDF bertambah mengikuti data yang telah dimasukkan oleh para Auditor.",
        "Urutan lingkup mengikuti urutan laporan baku, bukan urutan waktu pengisian atau nama pengisi.",
        "Kolom tanda tangan pada laporan disediakan kosong untuk penandatanganan manual.",
    ])
    add_page_break(doc)
    add_heading(doc, "Urutan lingkup pada laporan", 2)
    add_numbered(doc, [
        "Audit Internal Mutu",
        "Audit Internal Teknik Kelistrikan Jakarta",
        "Audit Internal Teknik Tekanan Jakarta",
        "Audit Internal Teknik Suhu dan Kelembapan Jakarta",
        "Audit Internal Teknik Vibrasi Jakarta",
        "Audit Internal Teknik Kelistrikan Gresik",
        "Audit Internal Teknik Tekanan Gresik",
        "Audit Internal Teknik Suhu Gresik",
        "Audit Internal Teknik Dimensi Gresik",
    ])

    add_page_break(doc)
    add_heading(doc, "10. Checklist Cepat Sebelum Selesai")
    add_heading(doc, "Checklist Auditor", 2)
    add_bullets(doc, [
        "Tahun laporan sudah benar.",
        "Peran Auditor dan lingkup tugas yang tampil sudah sesuai.",
        "Pakta Integritas sudah disimpan dan tanda tangan tersimpan.",
        "Klausul, kategori, dan uraian temuan sudah benar.",
        "Temuan sudah muncul pada tabel dan ringkasan jumlah sudah bertambah.",
    ])
    add_heading(doc, "Checklist Auditee", 2)
    add_bullets(doc, [
        "Kartu Mode Auditee | Lihat Saja sudah muncul pada bagian bawah.",
        "Temuan yang ditujukan kepada lingkup Auditee dapat dibaca.",
        "Tidak ada kebutuhan untuk menambah atau mengubah temuan dari mode Auditee.",
    ])
    add_heading(doc, "Jika tombol atau data belum muncul", 2)
    add_numbered(doc, [
        "Periksa tahun pada Periode Laporan.",
        "Klik Muat Ulang.",
        "Pastikan mapping akun sudah dibuat pada tahun tersebut.",
        "Untuk Cetak Pakta, pastikan Pakta sudah disimpan.",
        "Untuk Cetak Laporan, pastikan minimal satu temuan sudah tersimpan pada tahun tersebut.",
    ])
    add_callout(doc, "RINGKAS", "Auditor mengisi Pakta dan temuan. Auditee membaca temuan pada lingkupnya. Cetak Pakta mencetak dokumen per orang, sedangkan Cetak Laporan menampilkan rekap gabungan seluruh lingkup dalam satu tahun.")

    props = doc.core_properties
    props.title = "Panduan Penggunaan Fitur Temuan Ketidaksesuaian"
    props.subject = "Panduan Auditor Internal, Auditee, Pakta Integritas, dan Cetak Laporan"
    props.author = "ULAB"
    props.keywords = "Temuan Ketidaksesuaian, Auditor, Auditee, Pakta Integritas, Mapping Auditor"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
