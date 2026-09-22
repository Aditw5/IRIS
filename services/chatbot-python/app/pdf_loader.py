import os
import re
from typing import List, Dict, Optional
from docx import Document


COMMON_PREFIXES = [
    "manual book introduction website u-lab umro",
    "manual book introduction website u lab umro",
    "website u-lab umro",
    "website u lab umro",
]


SKIP_PREFIXES = [
    # Dihapus agar jawaban chatbot tidak mengambil pola jawaban singkat
    "jawaban singkat:",
    "jawaban singkat untuk chatbot:",
]


MAIN_HEADING_KEYWORDS = [
    "CARA",
    "TUJUAN",
    "AKSES",
    "ARTI",
    "MASALAH",
    "KATA KUNCI",
    "DAFTAR PERTANYAAN",
]


STEP_START_WORDS = [
    "buka",
    "klik",
    "isi",
    "pilih",
    "upload",
    "unggah",
    "masukkan",
    "lengkapi",
    "periksa",
    "pastikan",
    "tunggu",
    "cetak",
    "scan",
    "simpan",
    "gunakan",
    "catat",
    "cari",
    "lanjutkan",
    "masuk",
]


def normalize_spaces(text: str) -> str:
    return " ".join(str(text or "").strip().split())


def clean_paragraph(text: str) -> str:
    cleaned = normalize_spaces(text)
    lowered = cleaned.lower()

    if not cleaned:
        return ""

    for prefix in COMMON_PREFIXES:
        if lowered.startswith(prefix):
            return ""

    for prefix in SKIP_PREFIXES:
        if lowered.startswith(prefix):
            return ""

    return cleaned


def get_upper_ratio(text: str) -> float:
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return 0.0

    upper_count = sum(1 for c in letters if c.isupper())
    return upper_count / len(letters)


def looks_like_step_number(text: str) -> bool:
    """
    Mencegah langkah kerja seperti:
    1. Buka Dashboard Registrasi.
    2. Klik Navigation.
    dianggap sebagai heading section.
    """

    cleaned = normalize_spaces(text)
    match = re.match(r"^\d+\.\s+(.+)$", cleaned)

    if not match:
        return False

    after_number = match.group(1).strip()
    first_word = after_number.split()[0].lower() if after_number.split() else ""

    if first_word in STEP_START_WORDS:
        return True

    # Jika setelah nomor bukan huruf kapital dominan, biasanya ini langkah biasa.
    # Contoh: "1. Buka Dashboard Registrasi." upper ratio rendah.
    if get_upper_ratio(after_number) < 0.55:
        return True

    return False


def is_heading(text: str) -> bool:
    """
    Deteksi heading utama dokumen manual.

    Yang dianggap heading:
    - 3. CARA LOGIN KE DASHBOARD REGISTRASI
    - 10. CARA CETAK SERTIFIKAT ATAU LAPORAN REPAIR
    - 20. MASALAH YANG SERING TERJADI DAN SOLUSINYA
    - 21. KATA KUNCI PENTING UNTUK CHATBOT

    Yang TIDAK dianggap heading:
    - 1. Buka Dashboard Registrasi.
    - 2. Klik Navigation.
    - 3. Cari menu Alat Unit.
    """

    cleaned = normalize_spaces(text)

    if not cleaned:
        return False

    # Heading bernomor: "6. CARA MENAMBAHKAN ALAT BARU KE MASTER DATA"
    numbered_match = re.match(r"^\d+\.\s+(.+)$", cleaned)
    if numbered_match:
        title = numbered_match.group(1).strip()
        title_upper = title.upper()

        # Jangan anggap langkah bernomor sebagai heading
        if looks_like_step_number(cleaned):
            return False

        # Heading utama biasanya diawali kata ini
        for keyword in MAIN_HEADING_KEYWORDS:
            if title_upper.startswith(keyword):
                return True

        # Fallback untuk dokumen lama: heading dominan huruf besar
        if get_upper_ratio(title) >= 0.70 and len(title) <= 140:
            return True

        return False

    # Heading tanpa nomor, misalnya: "KATA KUNCI PENTING UNTUK CHATBOT"
    upper_ratio = get_upper_ratio(cleaned)
    if upper_ratio >= 0.75 and len(cleaned) <= 140:
        return True

    return False


def extract_text_from_docx(docx_path: str) -> str:
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"File DOCX tidak ditemukan: {docx_path}")

    document = Document(docx_path)
    paragraphs = []

    for paragraph in document.paragraphs:
        text = clean_paragraph(paragraph.text)
        if text:
            paragraphs.append(text)

    return "\n".join(paragraphs)


def list_docx_files(manuals_dir: str) -> List[str]:
    if not os.path.exists(manuals_dir):
        return []

    return sorted(
        [
            file_name for file_name in os.listdir(manuals_dir)
            if file_name.lower().endswith(".docx") and not file_name.startswith("~$")
        ]
    )


def extract_sections_from_docx(docx_path: str) -> List[Dict]:
    """
    Pecah dokumen DOCX menjadi section berdasarkan heading utama.

    Perbaikan penting:
    - Langkah bernomor seperti "1. Buka Dashboard" tidak lagi dianggap heading.
    - Paragraf "Jawaban singkat:" dibuang agar chatbot tidak memakai jawaban pendek.
    """

    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"File DOCX tidak ditemukan: {docx_path}")

    document = Document(docx_path)

    sections: List[Dict] = []
    current_heading: Optional[str] = None
    current_lines: List[str] = []

    for paragraph in document.paragraphs:
        text = clean_paragraph(paragraph.text)
        if not text:
            continue

        if is_heading(text):
            if current_heading and current_lines:
                sections.append({
                    "heading": current_heading,
                    "content": "\n".join(current_lines).strip()
                })

            current_heading = text
            current_lines = []
            continue

        if current_heading is None:
            # Abaikan intro sebelum heading pertama
            continue

        current_lines.append(text)

    if current_heading and current_lines:
        sections.append({
            "heading": current_heading,
            "content": "\n".join(current_lines).strip()
        })

    return sections


def load_all_manual_chunks(manuals_dir: str) -> List[Dict]:
    """
    Load semua manual DOCX menjadi chunk berbasis section.

    Tidak mengubah logika role customer/pelaksana/manager.
    File tetap dibaca berdasarkan semua DOCX di folder manuals.
    """

    if not os.path.exists(manuals_dir):
        raise FileNotFoundError(f"Folder manuals tidak ditemukan: {manuals_dir}")

    all_chunks: List[Dict] = []
    chunk_id = 1

    for file_name in list_docx_files(manuals_dir):
        docx_path = os.path.join(manuals_dir, file_name)

        try:
            sections = extract_sections_from_docx(docx_path)

            for index, section in enumerate(sections, start=1):
                all_chunks.append({
                    "id": chunk_id,
                    "file_name": file_name,
                    "chunk_index": index,
                    "heading": section["heading"],
                    "content": section["content"],
                })
                chunk_id += 1

        except Exception as e:
            all_chunks.append({
                "id": chunk_id,
                "file_name": file_name,
                "chunk_index": 0,
                "heading": "ERROR",
                "content": f"Gagal membaca file {file_name}: {str(e)}",
            })
            chunk_id += 1

    return all_chunks